from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any


READABLE_MCP_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".log",
    ".json",
    ".csv",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".html",
    ".css",
    ".java",
    ".sql",
}

BLOCKED_MCP_FILE_EXTENSIONS = {
    ".7z",
    ".bat",
    ".cmd",
    ".db",
    ".dll",
    ".env",
    ".exe",
    ".key",
    ".pem",
    ".ps1",
    ".rar",
    ".sh",
    ".sqlite",
    ".zip",
}

BLOCKED_MCP_FILE_NAMES = {".env", ".npmrc", ".pypirc"}


@dataclass(frozen=True)
class McpReadResult:
    text: str
    content_type: str
    visual_extraction_status: str = "not_required"
    warnings: tuple[str, ...] = ()


class McpFileReaderService:
    def can_read(self, path: Path) -> bool:
        return path.suffix.lower() in READABLE_MCP_EXTENSIONS and not self.is_blocked(path)

    def is_blocked(self, path: Path) -> bool:
        return path.name.lower() in BLOCKED_MCP_FILE_NAMES or path.suffix.lower() in BLOCKED_MCP_FILE_EXTENSIONS

    def read_text(self, path: Path, max_chars: int = 30_000) -> McpReadResult:
        extension = path.suffix.lower()
        if self.is_blocked(path):
            raise ValueError("Blocked file type.")
        if extension not in READABLE_MCP_EXTENSIONS:
            raise ValueError("This file type is not supported for readable indexing.")
        if extension == ".pdf":
            return self._read_pdf(path, max_chars)
        if extension == ".docx":
            return self._read_docx(path, max_chars)
        if extension == ".pptx":
            return self._read_pptx(path, max_chars)
        return self._read_plain_text(path, max_chars)

    def _read_plain_text(self, path: Path, max_chars: int) -> McpReadResult:
        text = path.read_text(encoding="utf-8", errors="replace")[:max_chars]
        return McpReadResult(text=text, content_type="text")

    def _read_pdf(self, path: Path, max_chars: int) -> McpReadResult:
        try:
            from pypdf import PdfReader
        except Exception as error:
            raise ValueError("PDF extraction dependency is not installed.") from error
        reader = PdfReader(str(path))
        chunks: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                chunks.append(f"--- Page {index} ---\n{text}")
            if sum(len(chunk) for chunk in chunks) >= max_chars:
                break
        combined = "\n\n".join(chunks).strip()[:max_chars]
        visual_status = "not_required" if combined else "needed"
        warnings = () if combined else ("No selectable text found. This PDF may be scanned/image-based.",)
        return McpReadResult(text=combined, content_type="pdf", visual_extraction_status=visual_status, warnings=warnings)

    def _read_docx(self, path: Path, max_chars: int) -> McpReadResult:
        try:
            from docx import Document
        except Exception as error:
            raise ValueError("DOCX extraction dependency is not installed.") from error
        document = Document(str(path))
        chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            rows: list[str] = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                chunks.append("\n".join(rows))
        text = "\n\n".join(chunks).strip()[:max_chars]
        if not text:
            raise ValueError("Could not extract text from DOCX.")
        return McpReadResult(text=text, content_type="docx")

    def _read_pptx(self, path: Path, max_chars: int) -> McpReadResult:
        try:
            from pptx import Presentation
        except Exception as error:
            raise ValueError("PPTX extraction dependency is not installed.") from error
        presentation = Presentation(str(path))
        chunks: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            slide_chunks = self._slide_text_chunks(slide)
            notes = self._slide_notes(slide)
            if notes:
                slide_chunks.append("Notes:\n" + notes)
            if slide_chunks:
                chunks.append(f"--- Slide {index} ---\n" + "\n\n".join(slide_chunks))
            if sum(len(chunk) for chunk in chunks) >= max_chars:
                break
        combined = "\n\n".join(chunks).strip()[:max_chars]
        visual_status = "not_required" if combined else "needed"
        warnings = () if combined else ("No text found in slides. This presentation may be image-heavy.",)
        return McpReadResult(text=combined, content_type="pptx", visual_extraction_status=visual_status, warnings=warnings)

    def _slide_text_chunks(self, slide: Any) -> list[str]:
        chunks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ).strip()
                if text:
                    chunks.append(text)
            if getattr(shape, "has_table", False):
                rows: list[str] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    chunks.append("Table:\n" + "\n".join(rows))
        return chunks

    def _slide_notes(self, slide: Any) -> str:
        try:
            notes_slide = slide.notes_slide
            frame = notes_slide.notes_text_frame
        except Exception:
            return ""
        return "\n".join(paragraph.text.strip() for paragraph in frame.paragraphs if paragraph.text.strip()).strip()


mcp_file_reader_service = McpFileReaderService()
